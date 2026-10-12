"""P1 / Part A — static-file extraction (Word/Excel/PDF/HTML → markdown + embedded images)."""
import base64
import io

import pytest

from integrations.extract import ExtractionError, extract


def _png_bytes() -> bytes:
    """A small real PNG, built with PyMuPDF (already a dependency)."""
    import pymupdf

    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 6, 6))
    pix.clear_with(180)
    return pix.tobytes("png")


def _docx_bytes(with_image: bool = False) -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading("Product Requirements", level=1)
    doc.add_paragraph("The system shall let a user log in.")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "ID"
    table.rows[0].cells[1].text = "Title"
    table.rows[1].cells[0].text = "US-1"
    table.rows[1].cells[1].text = "Login"
    if with_image:
        doc.add_picture(io.BytesIO(_png_bytes()))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _xlsx_bytes() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Stories"
    ws.append(["ID", "Title"])
    ws.append(["US-1", "Login"])
    ws.append(["US-2", "Logout"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pdf_bytes(text: str = "System Architecture\nExecutive Summary\nPattern: SPA",
               with_image: bool = False, blank: bool = False) -> bytes:
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    if not blank:
        page.insert_text((72, 90), text, fontsize=11)
    if with_image:
        page.insert_image(pymupdf.Rect(72, 200, 172, 300), stream=_png_bytes())
    return doc.tobytes()


def test_extract_docx_headings_paragraphs_tables():
    title, body, images = extract(filename="PRD.docx", mime="", content=_docx_bytes())
    assert title == "PRD"
    assert "# Product Requirements" in body
    assert "The system shall let a user log in." in body
    assert "| ID | Title |" in body and "| US-1 | Login |" in body
    assert images == []


def test_extract_docx_captures_embedded_image():
    _title, _body, images = extract(
        filename="spec.docx", mime="", content=_docx_bytes(with_image=True))
    assert len(images) == 1
    assert images[0]["bytes"] and images[0]["mime"].startswith("image/")


def test_extract_xlsx_sheets_as_tables():
    title, body, _images = extract(filename="backlog.xlsx", mime="", content=_xlsx_bytes())
    assert title == "backlog"
    assert "## Stories" in body
    assert "| ID | Title |" in body
    assert "US-1" in body and "US-2" in body


def test_extract_txt_returns_body_as_is():
    title, body, images = extract(
        filename="context.txt", mime="text/plain",
        content=b"# Notes\n\nThe billing service handles invoices.\n",
    )
    assert title == "context"
    assert "# Notes" in body
    assert "The billing service handles invoices." in body
    assert images == []


def test_extract_html_strips_tags_keeps_structure_and_images():
    data_uri = "data:image/png;base64," + base64.b64encode(_png_bytes()).decode()
    html = (
        "<html><head><style>.x{color:red}</style><title>t</title></head>"
        "<body><h1>Design</h1><p>Screens for <b>login</b>.</p>"
        "<ul><li>Home</li><li>Profile</li></ul>"
        f'<img src="{data_uri}" alt="hero"/>'
        "<script>alert('x')</script></body></html>"
    ).encode()
    title, body, images = extract(filename="ux.html", mime="text/html", content=html)
    assert title == "ux"
    assert "# Design" in body
    assert "Screens for login." in body
    assert "- Home" in body and "- Profile" in body
    assert "alert" not in body and "color:red" not in body
    # the data-URI image is captured with real bytes
    assert len(images) == 1 and images[0]["bytes"]


def test_extract_pdf_text():
    """pymupdf4llm yields clean markdown text (no injected 'Page N' headers)."""
    _t, body, _i = extract(filename="arch.pdf", mime="application/pdf", content=_pdf_bytes())
    assert "System Architecture" in body and "Pattern: SPA" in body
    assert "## Page" not in body


def test_extract_pdf_captures_embedded_image():
    _t, _body, images = extract(
        filename="wire.pdf", mime="application/pdf", content=_pdf_bytes(with_image=True))
    assert len(images) >= 1
    assert images[0]["bytes"] and images[0]["page"] == 1


def test_extract_pdf_empty_when_no_text_or_image():
    """An image-only/blank PDF yields no text — ingest turns '' into a helpful note (BF-BUG-3b)."""
    _t, body, images = extract(
        filename="scan.pdf", mime="application/pdf", content=_pdf_bytes(blank=True))
    assert body == "" and images == []


def test_pdf_page_pngs_one_per_page():
    """A2: each PDF page renders to a PNG (the faithful visual for the viewer + vision step)."""
    import pymupdf

    from integrations.extract import pdf_page_pngs

    doc = pymupdf.open()
    doc.new_page().insert_text((72, 90), "Page one")
    doc.new_page().insert_text((72, 90), "Page two")
    pngs = pdf_page_pngs(doc.tobytes())
    assert [p["page"] for p in pngs] == [1, 2]
    assert all(p["mime"] == "image/png" and p["bytes"][:8] == b"\x89PNG\r\n\x1a\n" for p in pngs)


def test_clean_pdf_markdown_flattens_br_and_drops_empty_tables():
    from integrations.extract import _clean_pdf_markdown

    # <br> (pymupdf4llm in-cell breaks) → space; a real table with content is kept.
    real = _clean_pdf_markdown("|A|B|\n|---|---|\n|x<br>y|z|")
    assert "<br>" not in real and "x y" in real and "|A|B|" in real
    # an image-only page's empty-table skeleton + bare heading → nothing (→ ingest shows a note).
    assert _clean_pdf_markdown("# \n\n||||\n|---|---|---|\n||||\n\n|||\n|---|---|\n|||") == ""


def test_unsupported_type_raises():
    with pytest.raises(ExtractionError):
        extract(filename="logo.png", mime="image/png", content=b"\x89PNG\r\n")


def test_corrupt_pdf_wrapped_as_extraction_error():
    with pytest.raises(ExtractionError):
        extract(filename="broken.pdf", mime="application/pdf", content=b"not a real pdf")
