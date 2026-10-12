"""P0-T11 — REST API surface.

Exercises the HTTP layer end-to-end against the real domain services, stubbing only
the Temporal-touching seam (`_default_workflow_starter` / `_default_signal_sender`) at
module level so Run/Approval rows are genuinely created. Tenant scoping is asserted
directly (the test connection is superuser, so RLS is bypassed — the explicit
`tenant_id` filter is what's under test here; RLS itself is covered by test_rls.py).
"""
import pytest
from rest_framework.test import APIClient

from apps.gates import services as gate_services
from apps.gates.models import Approval
from apps.projects.models import Project, ProjectProfile
from apps.registry.catalog import seed_agent_specs
from apps.registry.models import AgentSpec, Artifact, TraceEdge
from apps.runs.models import Run
from apps.tenancy.models import Membership, Tenant, User
from apps.workitems.models import WorkItem
from orchestration import trigger as trigger_mod


@pytest.fixture
def client():
    return APIClient()


def _hdr(tenant_id):
    return {"HTTP_X_TENANT_ID": str(tenant_id)}


@pytest.fixture
def tenant(db):
    return Tenant.objects.create(name="Acme")


@pytest.fixture
def project(tenant):
    return Project.objects.create(tenant=tenant, name="Platform")


@pytest.fixture
def lead(client, tenant):
    """A signed-in tenant admin — request this in any test that WRITES.

    Bugs S.No 214 closed the write endpoints that were still running under DRF's `AllowAny`
    default: creating projects and work items, editing the project profile, triggering a phase,
    uploading a source document, reading traceability, wiring Atlassian. The tests that exercise
    them were posting with no session at all, which is precisely the hole that was fixed, so each
    one now has to act as somebody. Tenant admin also satisfies the per-project role check on the
    two spend/steer endpoints (profile, trigger) via the same admin fallback
    `require_repo_manage_role` has always used."""
    user = User.objects.create_user(email="lead@acme.test", name="Lead", tenant=tenant,
                                    password="pw", is_tenant_admin=True)
    client.force_authenticate(user)
    return user


# ── Tenant bootstrap ────────────────────────────────────────────────
def test_create_tenant(client, db):
    resp = client.post("/api/tenants/", {"name": "NewCo"}, format="json")
    assert resp.status_code == 201
    assert resp.data["name"] == "NewCo"
    tenant = Tenant.objects.get(id=resp.data["id"], name="NewCo")
    assert Tenant.objects.filter(id=tenant.id, name="NewCo").exists()
    assert AgentSpec.objects.filter(tenant=tenant).count() > 0
    assert AgentSpec.objects.filter(tenant=tenant, key="01_ideation").exists()


# ── Projects + profile ──────────────────────────────────────────────
def test_create_and_list_projects(client, tenant, lead):
    resp = client.post("/api/projects/", {"name": "Apollo"}, format="json", **_hdr(tenant.id))
    assert resp.status_code == 201
    pid = resp.data["id"]
    assert Project.objects.get(id=pid).tenant_id == tenant.id

    listed = client.get("/api/projects/", **_hdr(tenant.id))
    assert listed.status_code == 200
    assert [p["id"] for p in listed.data] == [pid]


def test_duplicate_project_name_is_refused_with_409(client, tenant, lead):
    """Bugs S.No 211: two projects could share a name, so the switcher, the board header and
    every invite named two different things identically. The Create Project screen has always
    had the collision state built and branches on 409 — the backend just never returned one,
    which is what made the state unreachable."""
    assert client.post("/api/projects/", {"name": "Billing Service"}, format="json",
                       **_hdr(tenant.id)).status_code == 201

    dup = client.post("/api/projects/", {"name": "Billing Service"}, format="json",
                      **_hdr(tenant.id))
    assert dup.status_code == 409
    assert "already exists" in dup.data["detail"]

    # Case-insensitive: "billing service" is the same project to whoever reads the switcher.
    cased = client.post("/api/projects/", {"name": "billing service"}, format="json",
                        **_hdr(tenant.id))
    assert cased.status_code == 409

    # ...and so is one that differs only by the spaces around it.
    padded = client.post("/api/projects/", {"name": "  Billing Service  "}, format="json",
                         **_hdr(tenant.id))
    assert padded.status_code == 409
    assert Project.objects.filter(tenant=tenant).count() == 1


def test_another_workspace_may_reuse_the_name(client, db, tenant, lead):
    """The name is unique per workspace, not globally — two customers both having a "Billing
    Service" is normal and must not be blocked."""
    assert client.post("/api/projects/", {"name": "Billing Service"}, format="json",
                       **_hdr(tenant.id)).status_code == 201

    other = Tenant.objects.create(name="Other Co")
    user = User.objects.create_user(email="lead@other.test", name="L", tenant=other,
                                    password="pw", is_tenant_admin=True)
    client.force_authenticate(user)
    assert client.post("/api/projects/", {"name": "Billing Service"}, format="json",
                       **_hdr(other.id)).status_code == 201


def test_blank_project_name_is_refused(client, tenant, lead):
    """A name of nothing but spaces was accepted and then rendered as an empty, unidentifiable
    switcher entry. The Create button already blocks this, so only a direct API call gets here."""
    for name in ("", "   ", "\t\n"):
        resp = client.post("/api/projects/", {"name": name}, format="json", **_hdr(tenant.id))
        assert resp.status_code == 400, f"{name!r} was accepted"
    assert Project.objects.filter(tenant=tenant).count() == 0


def test_project_name_is_stored_trimmed(client, tenant, lead):
    resp = client.post("/api/projects/", {"name": "  Apollo  "}, format="json", **_hdr(tenant.id))
    assert resp.status_code == 201
    assert Project.objects.get(id=resp.data["id"]).name == "Apollo"


def test_profile_upsert_bumps_version(client, project, lead):
    url = f"/api/projects/{project.id}/profile/"
    # GET before any profile → 404
    assert client.get(url, **_hdr(project.tenant_id)).status_code == 404

    created = client.put(
        url, {"important_context": "north-star"}, format="json", **_hdr(project.tenant_id)
    )
    assert created.status_code == 201
    assert created.data["important_context"] == "north-star"
    assert created.data["version"] == 1

    updated = client.put(
        url, {"important_context": "revised"}, format="json", **_hdr(project.tenant_id)
    )
    assert updated.status_code == 200
    assert updated.data["version"] == 2
    assert ProjectProfile.objects.get(project=project).important_context == "revised"


# ── WorkItems ───────────────────────────────────────────────────────
def test_create_and_filter_work_items(client, project, lead):
    resp = client.post(
        "/api/work-items/",
        {"project": str(project.id), "type": "story", "feature_stream": "auth"},
        format="json", **_hdr(project.tenant_id),
    )
    assert resp.status_code == 201
    wid = resp.data["id"]
    assert WorkItem.objects.get(id=wid).state == WorkItem.State.PROPOSED

    by_project = client.get(f"/api/work-items/?project={project.id}", **_hdr(project.tenant_id))
    assert [w["id"] for w in by_project.data] == [wid]
    # State filter that excludes it returns empty.
    none = client.get("/api/work-items/?state=done", **_hdr(project.tenant_id))
    assert none.data == []


def test_create_work_item_rejects_foreign_project(client, tenant, lead):
    other = Tenant.objects.create(name="Other")
    foreign = Project.objects.create(tenant=other, name="Theirs")
    resp = client.post(
        "/api/work-items/", {"project": str(foreign.id)}, format="json", **_hdr(tenant.id)
    )
    assert resp.status_code == 400
    assert "project" in resp.data


def test_patch_work_item_repo_assign_reject_foreign_and_clear(client, project, lead):
    """B1 (Option C): a work item's target repo is assignable/clearable via PATCH and validated
    to a repository on its OWN project — this backs the reviewer override + per-task redirect."""
    from apps.projects.models import Repository
    wi = WorkItem.objects.create(tenant=project.tenant, project=project, type=WorkItem.Type.TASK)
    repo = Repository.objects.create(tenant=project.tenant, project=project, name="acme/api",
                                     role="backend")
    ok = client.patch(f"/api/work-items/{wi.id}/", {"repo": str(repo.id)},
                      format="json", **_hdr(project.tenant_id))
    assert ok.status_code == 200
    wi.refresh_from_db()
    assert str(wi.repo_id) == str(repo.id)

    # a repo on a DIFFERENT project (same tenant) is rejected
    other = Project.objects.create(tenant=project.tenant, name="Other")
    foreign = Repository.objects.create(tenant=project.tenant, project=other, name="x/y",
                                        role="backend")
    bad = client.patch(f"/api/work-items/{wi.id}/", {"repo": str(foreign.id)},
                       format="json", **_hdr(project.tenant_id))
    assert bad.status_code == 400 and "repo" in bad.data

    # clearing the assignment (null) is allowed — falls back to the default repo at build
    clr = client.patch(f"/api/work-items/{wi.id}/", {"repo": None},
                       format="json", **_hdr(project.tenant_id))
    assert clr.status_code == 200
    wi.refresh_from_db()
    assert wi.repo_id is None


# ── Phase trigger (real trigger_phase, stubbed workflow start) ──────
def test_trigger_phase_creates_run(client, project, monkeypatch, lead):
    seed_agent_specs(project.tenant)
    wi = WorkItem.objects.create(tenant=project.tenant, project=project, type=WorkItem.Type.STORY)
    started = []
    monkeypatch.setattr(
        trigger_mod, "_default_workflow_starter",
        lambda wf_id, inp: started.append((wf_id, inp.phase)),
    )

    resp = client.post(
        f"/api/work-items/{wi.id}/trigger/", {"agent_key": "01_ideation"},
        format="json", **_hdr(project.tenant_id),
    )
    assert resp.status_code == 202
    run = Run.objects.get(id=resp.data["run_id"])
    assert run.work_item_id == wi.id
    assert run.temporal_workflow_id == resp.data["workflow_id"]
    assert started and started[0][1] == "ideation"


def test_trigger_unknown_agent_is_400(client, project, monkeypatch, lead):
    wi = WorkItem.objects.create(tenant=project.tenant, project=project, type=WorkItem.Type.STORY)
    monkeypatch.setattr(trigger_mod, "_default_workflow_starter", lambda *a, **k: None)
    resp = client.post(
        f"/api/work-items/{wi.id}/trigger/", {"agent_key": "nope"},
        format="json", **_hdr(project.tenant_id),
    )
    assert resp.status_code == 400


# ── Runs / artifacts / events / trace-edges (read) ──────────────────
def test_list_runs_artifacts_events_edges(client, project):
    spec = AgentSpec.objects.create(tenant=project.tenant, key="03_prd", body_ref="x")
    wi = WorkItem.objects.create(tenant=project.tenant, project=project, type=WorkItem.Type.STORY)
    run = Run.objects.create(tenant=project.tenant, work_item=wi, agent_spec=spec)
    a1 = Artifact.objects.create(
        tenant=project.tenant, project=project, run=run, phase="requirements", type="requirements"
    )
    a2 = Artifact.objects.create(
        tenant=project.tenant, project=project, run=run, phase="prd", type="prd"
    )
    TraceEdge.objects.create(tenant=project.tenant, from_artifact=a1, to_artifact=a2, rel="derives")

    runs = client.get(f"/api/runs/?project={project.id}", **_hdr(project.tenant_id))
    assert [r["id"] for r in runs.data] == [str(run.id)]

    arts = client.get(f"/api/artifacts/?project={project.id}&phase=prd", **_hdr(project.tenant_id))
    assert [a["id"] for a in arts.data] == [str(a2.id)]

    edges = client.get(f"/api/trace-edges/?project={project.id}", **_hdr(project.tenant_id))
    assert edges.data[0]["rel"] == "derives"

    events = client.get(f"/api/events/?target_id={run.id}", **_hdr(project.tenant_id))
    assert events.status_code == 200


def test_traceability_endpoint_returns_nodes_and_edges(client, project, lead):
    run = Run.objects.create(
        tenant=project.tenant, work_item=WorkItem.objects.create(
            tenant=project.tenant, project=project, type=WorkItem.Type.STORY),
        agent_spec=AgentSpec.objects.create(tenant=project.tenant, key="03_prd", body_ref="x"),
    )
    req = Artifact.objects.create(
        tenant=project.tenant, project=project, run=run, phase="requirements",
        type="requirements", status="published", sor="internal",
    )
    prd = Artifact.objects.create(
        tenant=project.tenant, project=project, run=run, phase="prd", type="prd",
        status="published", sor="confluence",
        external_ref={"sor": "confluence", "id": "123", "url": "https://x/wiki/123", "version": 1},
    )
    TraceEdge.objects.create(
        tenant=project.tenant, from_artifact=req, to_artifact=prd, rel="derives"
    )

    resp = client.get(f"/api/projects/{project.id}/traceability/", **_hdr(project.tenant_id))
    assert resp.status_code == 200
    nodes = {n["id"]: n for n in resp.data["nodes"]}
    assert set(nodes) == {str(req.id), str(prd.id)}
    # The PRD node carries its live SoR deep-link.
    assert nodes[str(prd.id)]["sor"] == "confluence"
    assert nodes[str(prd.id)]["url"] == "https://x/wiki/123"
    assert nodes[str(req.id)]["url"] is None
    assert resp.data["edges"] == [
        {"from": str(req.id), "to": str(prd.id), "rel": "derives"}
    ]


# ── Gate decision (real decide_gate, stubbed signal) ────────────────
@pytest.fixture
def gate(project, monkeypatch):
    spec = AgentSpec.objects.create(tenant=project.tenant, key="01_ideation", body_ref="x")
    wi = WorkItem.objects.create(tenant=project.tenant, project=project, type=WorkItem.Type.STORY)
    run = Run.objects.create(
        tenant=project.tenant, work_item=wi, agent_spec=spec, temporal_workflow_id="wf-x"
    )
    t = project.tenant
    po = User.objects.create_user(email="po@acme.test", name="PO", tenant=t, password="p")
    Membership.objects.create(tenant=t, user=po, project=project, role="product_owner")
    eng = User.objects.create_user(email="e@acme.test", name="E", tenant=t, password="p")
    Membership.objects.create(tenant=t, user=eng, project=project, role="engineer")
    approval = Approval.objects.create(
        tenant=project.tenant, run=run, work_item=wi, required_role="product_owner",
    )
    signals = []
    monkeypatch.setattr(
        gate_services, "_default_signal_sender",
        # `revise` is part of the signal contract (a rejection can ask for a revision
        # rather than end the run); a stub that omits it fails on the real call.
        lambda wf, dec, uid, comment, revise=False: signals.append((wf, dec)),
    )
    return {"approval": approval, "po": po, "eng": eng, "signals": signals}


def test_decide_gate_approved(client, project, gate):
    resp = client.post(
        f"/api/approvals/{gate['approval'].id}/decide/",
        {"user_id": str(gate["po"].id), "decision": "approved", "comment": "lgtm"},
        format="json", **_hdr(project.tenant_id),
    )
    assert resp.status_code == 200
    gate["approval"].refresh_from_db()
    assert gate["approval"].decision == Approval.Decision.APPROVED
    assert gate["signals"] == [("wf-x", "approved")]


def test_decide_gate_wrong_role_is_403(client, project, gate):
    resp = client.post(
        f"/api/approvals/{gate['approval'].id}/decide/",
        {"user_id": str(gate["eng"].id), "decision": "approved"},
        format="json", **_hdr(project.tenant_id),
    )
    assert resp.status_code == 403
    gate["approval"].refresh_from_db()
    assert gate["approval"].decision == Approval.Decision.PENDING
    assert gate["signals"] == []


def test_list_pending_approvals(client, project, gate):
    resp = client.get("/api/approvals/?decision=pending", **_hdr(project.tenant_id))
    assert [a["id"] for a in resp.data] == [str(gate["approval"].id)]


# ── Brownfield-lite onboarding ──────────────────────────────────────
def _docx_bytes():
    import io

    from docx import Document

    doc = Document()
    doc.add_heading("PRD", level=1)
    doc.add_paragraph("The system shall let a user log in.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_source_document_upload_extract_and_adopt(client, project, monkeypatch, lead):
    from django.core.files.uploadedfile import SimpleUploadedFile

    from integrations import onboarding
    from storage import InMemoryObjectStore

    monkeypatch.setattr(onboarding, "_object_store", lambda injected: InMemoryObjectStore())
    po = User.objects.create_user(email="po@x.test", name="PO", tenant=project.tenant, password="p")
    Membership.objects.create(tenant=project.tenant, user=po, project=project, role="product_owner")

    upload = SimpleUploadedFile(
        "PRD.docx", _docx_bytes(),
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    up = client.post(
        f"/api/projects/{project.id}/source-documents/",
        {"file": upload, "artifact_type": "prd"}, format="multipart", **_hdr(project.tenant_id),
    )
    assert up.status_code == 201
    sd_id = up.data["source_document_id"]
    # An upload extracts but registers NO artifact — the registry entry is created once, at
    # adopt, so a document uploaded once is one artifact at v1 rather than a v1 the upload
    # made and a v2 that replaced it. Everything shown before adopting comes from the
    # source document itself.
    assert up.data["artifact_ids"] == []
    assert not Artifact.objects.filter(project=project).exists()

    listed = client.get(
        f"/api/projects/{project.id}/source-documents/", **_hdr(project.tenant_id)
    )
    assert [s["id"] for s in listed.data] == [sd_id]

    # Review gate: product_owner adopts → the artifact is created, published.
    res = client.post(
        f"/api/source-documents/{sd_id}/adopt/", {"user_id": str(po.id)},
        format="json", **_hdr(project.tenant_id),
    )
    assert res.status_code == 200
    adopted = Artifact.objects.filter(project=project)
    assert adopted.count() == 1
    assert adopted.first().status == "published"


def test_onboard_atlassian_endpoint(client, project, monkeypatch, lead):
    from api import views

    monkeypatch.setattr(
        views, "register_atlassian",
        lambda **kw: {"registered_pages": ["a1"], "registered_issues": ["ST-1"]},
    )
    res = client.post(
        f"/api/projects/{project.id}/onboard-atlassian/", {}, format="json",
        **_hdr(project.tenant_id),
    )
    assert res.status_code == 200
    assert res.data == {"registered_pages": ["a1"], "registered_issues": ["ST-1"]}


# ── Tenant scoping ──────────────────────────────────────────────────
def test_tenant_isolation_hides_other_tenants_rows(client, project):
    other = Tenant.objects.create(name="Other")
    listed = client.get("/api/projects/", **_hdr(other.id))
    assert listed.data == []
    # Direct retrieve of a foreign project is a 404 under the other tenant.
    retr = client.get(f"/api/projects/{project.id}/", **_hdr(other.id))
    assert retr.status_code == 404


def test_missing_tenant_header_is_400(client):
    assert client.get("/api/projects/").status_code == 400


# ── S.No 214: the write endpoints that were running under AllowAny ──
def _anon_writes(project):
    """(label, method, path, body) for every write this fix closed."""
    pid = project.id
    return [
        ("create project", "post", "/api/projects/", {"name": "Sneaky"}),
        ("edit profile", "put", f"/api/projects/{pid}/profile/", {"domain": "x"}),
        ("upload source doc", "post", f"/api/projects/{pid}/source-documents/", {}),
        ("wire atlassian", "post", f"/api/projects/{pid}/onboard-atlassian/", {}),
        ("create work item", "post", "/api/work-items/",
         {"project": str(pid), "type": "story", "title": "t"}),
    ]


def test_the_write_endpoints_refuse_an_anonymous_caller(client, project):
    """Bugs S.No 214. DRF's project-wide default is `AllowAny`, and these actions never stated
    their own permission, so anyone who could reach the API could create projects, rewrite the
    project profile every agent reads as its briefing, create work items and mint a companion
    credential — with no login at all. Tenant scoping and RLS still applied, so this was never
    cross-tenant exposure; it was that INSIDE a tenant nothing checked who you were."""
    hdr = _hdr(project.tenant_id)
    failures = [
        f"{label} → "
        f"{getattr(client, method)(path, body, format='json', **hdr).status_code}"
        for label, method, path, body in _anon_writes(project)
        if getattr(client, method)(path, body, format="json",
                                   **hdr).status_code not in (401, 403)
    ]
    assert not failures, "these still accept an anonymous caller: " + "; ".join(failures)


def test_a_signed_in_lead_can_still_do_all_of_it(client, project, lead):
    """The other half: the fix must not have closed the door on the people who belong here.
    Anything but 401/403 is a pass — several of these legitimately 400 on a stub body; what is
    under test is that the permission layer let them through."""
    blocked = [
        label for label, method, path, body in _anon_writes(project)
        if getattr(client, method)(path, body, format="json",
                                   **_hdr(project.tenant_id)).status_code in (401, 403)
    ]
    assert not blocked, f"a signed-in tenant admin was refused: {blocked}"


def test_the_second_layer_is_scoped_to_what_it_is_actually_for(client, project):
    """Layer 2, and the line it draws. `IsAuthenticatedMember` only asks 'is anybody logged in' —
    its own docstring says so — which inside a tenant still lets a signed-in user act on a project
    they are not on. Two different bars fix that without breaking the journey:

    * the **project profile** takes tech_lead/admin, because rewriting it re-steers every agent
      that will ever run on the project;
    * everything else takes **membership**, because each role drives its own phase — a Product
      Owner triggers ideation, an Engineer triggers implementation, a developer connects their own
      laptop. A tech_lead-only bar there would have locked most of the team out of ordinary work,
      which is why it is not the bar used.
    """
    member = User.objects.create_user(email="dev@acme.test", name="Dev",
                                      tenant=project.tenant, password="pw")
    Membership.objects.create(tenant=project.tenant, user=member, project=project,
                              role="engineer")
    client.force_authenticate(member)
    wi = WorkItem.objects.create(tenant=project.tenant, project=project, type="story", title="t")
    hdr = _hdr(project.tenant_id)

    # Steering every agent on the project is not an ordinary member's call.
    assert client.put(f"/api/projects/{project.id}/profile/", {"domain": "x"},
                      format="json", **hdr).status_code == 403
    # Their own everyday work is untouched.
    made = client.post("/api/work-items/",
                       {"project": str(project.id), "type": "story", "title": "mine"},
                       format="json", **hdr)
    assert made.status_code == 201
    # Triggering reaches the domain layer rather than being refused at the permission layer.
    assert client.post(f"/api/work-items/{wi.id}/trigger/", {"agent_key": "03_prd"},
                       format="json", **hdr).status_code != 403


def test_a_signed_in_outsider_cannot_act_on_a_project_they_are_not_on(client, project):
    """The hole membership closes: a session is not the same as belonging here."""
    outsider = User.objects.create_user(email="other@acme.test", name="Other",
                                        tenant=project.tenant, password="pw")
    client.force_authenticate(outsider)
    wi = WorkItem.objects.create(tenant=project.tenant, project=project, type="story", title="t")
    assert client.post(f"/api/work-items/{wi.id}/trigger/", {"agent_key": "03_prd"},
                       format="json", **_hdr(project.tenant_id)).status_code == 403


# ── N2: a link's text is a promise — the API has to be able to keep it ──
def test_an_approval_carries_BOTH_vocabularies(client, project):
    """The vocabulary trap, pinned. `Approval.phase` is the agent KEY and the UI's phase-specific
    gate panels key on it; `Artifact.phase` is the catalog's phase LABEL. They differ, so
    `/artifacts?phase=<approval.phase>` returns ZERO rows — a silent empty list, worse than the
    wrong-but-populated one it would replace.

    `00c` is the case that breaks the obvious client-side workaround (strip the `NN_` prefix):
    that yields `intelligent_context_extractor`, and the real label is `context_extraction`. It
    is right for 10 of the 11 agents, which is exactly what makes it dangerous."""
    spec = AgentSpec.objects.create(tenant=project.tenant, key="00c_intelligent_context_extractor",
                                    body_ref="x", phase_label="context_extraction")
    wi = WorkItem.objects.create(tenant=project.tenant, project=project, type=WorkItem.Type.STORY)
    run = Run.objects.create(tenant=project.tenant, work_item=wi, agent_spec=spec,
                             temporal_workflow_id="wf-x")
    Approval.objects.create(tenant=project.tenant, run=run, work_item=wi,
                            required_role="product_owner")
    row = client.get("/api/approvals/?decision=pending", **_hdr(project.tenant_id)).json()[0]
    assert row["phase"] == "00c_intelligent_context_extractor"        # panels depend on this
    assert row["phase_label"] == "context_extraction"                 # artifacts filter on this
    assert row["phase"].split("_", 1)[1] != row["phase_label"]        # the heuristic would fail


@pytest.fixture
def two_runs_of_artifacts(project):
    spec = AgentSpec.objects.create(tenant=project.tenant, key="03_prd", body_ref="x",
                                    phase_label="prd")
    wi = WorkItem.objects.create(tenant=project.tenant, project=project, type=WorkItem.Type.STORY)
    runs = [Run.objects.create(tenant=project.tenant, work_item=wi, agent_spec=spec,
                               temporal_workflow_id=f"wf-{n}") for n in (1, 2)]
    made = []
    for n, run in enumerate(runs, start=1):
        made.append(Artifact.objects.create(
            tenant=project.tenant, project=project, run=run, type="prd", phase="prd",
            title="PRD", version=n, status=Artifact.Status.PUBLISHED, sor="internal"))
    made.append(Artifact.objects.create(
        tenant=project.tenant, project=project, run=runs[1], type="wireframe", phase="ux_design",
        title="Sign-in", version=1, status=Artifact.Status.DRAFT, sor="internal"))
    return {"runs": runs, "artifacts": made}


def test_artifacts_can_be_scoped_to_one_run(client, project, two_runs_of_artifacts):
    """The scope a gate actually has. It is the right one BECAUSE it does not require the platform
    to invent a notion of a run's *primary* artifact — the second run here registered two, and a
    "pick the first" rule would be a new lie in place of the old one."""
    run = two_runs_of_artifacts["runs"][1]
    rows = client.get(f"/api/artifacts/?run={run.id}", **_hdr(project.tenant_id)).json()
    assert {r["type"] for r in rows} == {"prd", "wireframe"}
    assert len(rows) == 2


def test_artifacts_can_be_scoped_to_one_type(client, project, two_runs_of_artifacts):
    rows = client.get("/api/artifacts/?type=wireframe", **_hdr(project.tenant_id)).json()
    assert [r["type"] for r in rows] == ["wireframe"]


def test_a_run_scope_turns_latest_off_so_a_superseded_version_still_resolves(
    client, project, two_runs_of_artifacts
):
    """The easiest thing in this whole change to get wrong. The registry hardcodes `latest`, which
    collapses to one row per (phase, type, title) — so a deep link into the run that produced v1
    would find NOTHING once a later run made v2. A run-scoped view is a PROVENANCE view; v1
    belongs in it."""
    first = two_runs_of_artifacts["runs"][0]
    rows = client.get(f"/api/artifacts/?run={first.id}&latest=1", **_hdr(project.tenant_id)).json()
    assert [r["version"] for r in rows] == [1], "the run's own v1 must still resolve"


def test_the_unscoped_registry_still_collapses_to_latest(client, project, two_runs_of_artifacts):
    """The other side of that rule — turning `latest` off must not leak into the plain registry."""
    rows = client.get("/api/artifacts/?latest=1", **_hdr(project.tenant_id)).json()
    prds = [r for r in rows if r["type"] == "prd"]
    assert [r["version"] for r in prds] == [2]


def test_the_existing_filters_behave_exactly_as_before(client, project, two_runs_of_artifacts):
    """The no-regression test for the `_FILTERS` refactor. These three are what the app already
    sends on every registry load."""
    tid = project.tenant_id
    assert len(client.get(f"/api/artifacts/?project={project.id}", **_hdr(tid)).json()) == 3
    assert len(client.get("/api/artifacts/?phase=prd", **_hdr(tid)).json()) == 2
    assert len(client.get("/api/artifacts/?status=draft", **_hdr(tid)).json()) == 1
    # An agent key is NOT a phase label, and must keep returning nothing rather than guessing.
    assert client.get("/api/artifacts/?phase=03_prd", **_hdr(tid)).json() == []


def test_artifacts_can_be_scoped_to_one_work_item(client, project, two_runs_of_artifacts):
    """An artifact has no direct work-item FK — it is reached through the run that produced it.
    That indirection is exactly why the count on a work-item card used to be inert text: there
    was no filter to link to."""
    run = two_runs_of_artifacts["runs"][0]
    wi = run.work_item_id
    rows = client.get(f"/api/artifacts/?work_item={wi}", **_hdr(project.tenant_id)).json()
    assert len(rows) == 3           # both runs belong to the same work item
    other = WorkItem.objects.create(tenant=project.tenant, project=project,
                                    type=WorkItem.Type.STORY, title="unrelated")
    assert client.get(f"/api/artifacts/?work_item={other.id}",
                      **_hdr(project.tenant_id)).json() == []
