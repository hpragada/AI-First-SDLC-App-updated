"""Build the QA execution tracker from the Markdown test-case suites.

    python docs/qa/build_tracker.py
    python docs/qa/build_tracker.py --out Greenfield_Test_Execution_Tracker_V1.xlsx

Parses every ``docs/qa/suites/TS-*.md`` file, extracts one row per test case, and
writes ``docs/qa/Greenfield_Test_Execution_Tracker.xlsx`` — or, with ``--out``, a new
workbook beside it under that name, leaving the original untouched. Logged results are
read from the target if it already exists, otherwise from the original, so a first
versioned build starts from everything already recorded.

Re-running is safe: execution results already logged in the workbook are carried
over by TC-ID, and the Defects and Daily Log sheets are preserved verbatim. New
test cases are appended as blank rows; a test case deleted from the Markdown
drops out of the Test Cases sheet but its defects stay.
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

QA_DIR = Path(__file__).resolve().parent
SUITES_DIR = QA_DIR / "suites"
OUT = QA_DIR / "Greenfield_Test_Execution_Tracker.xlsx"

# Suite -> the tester who owns it (see GREENFIELD_TEST_PLAN.md §5).
SUITE_OWNER = {
    "TS-01": "A", "TS-02": "A", "TS-03": "A",
    "TS-04": "B", "TS-06": "B", "TS-11": "B", "TS-12": "B", "TS-13": "B",
    "TS-05": "C", "TS-07": "C", "TS-08": "C", "TS-09": "C",
    "TS-10": "C", "TS-14": "C", "TS-15": "C",
}

STATUSES = ["Not Run", "Pass", "Fail", "Blocked", "Skipped"]
SEVERITIES = ["Critical", "High", "Medium", "Low"]
DEFECT_STATES = ["Open", "In Progress", "Fixed — awaiting re-test", "Closed", "Won't Fix", "Duplicate"]

# ── styling ────────────────────────────────────────────────────────────────
NAVY = "1B2A4A"
ORANGE = "E8743B"
LIGHT = "F2F4F8"
HEAD = Font(bold=True, color="FFFFFF", size=11)
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
TITLE = Font(bold=True, size=14, color=NAVY)
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")
THIN = Side(style="thin", color="D0D5DD")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

CASE_RE = re.compile(r"^###\s+(TC-\d+)\s*[·\-–]\s*(.+?)\s*$")
META_RE = re.compile(r"^\|\s*([^|]+?)\s*\|\s*(.*?)\s*\|\s*$")


def clean(text: str) -> str:
    """Strip the Markdown emphasis and code ticks used in metadata cells."""
    text = text.strip()
    text = re.sub(r"`([^`]*)`", r"\1", text)
    text = re.sub(r"\*\*([^*]*)\*\*", r"\1", text)
    return text.strip()


def parse_suites() -> list[dict]:
    """Return one dict per test case, in suite then TC-ID order."""
    if not SUITES_DIR.is_dir():
        sys.exit(f"No suites directory at {SUITES_DIR}")

    cases: list[dict] = []
    seen: dict[str, str] = {}

    for path in sorted(SUITES_DIR.glob("TS-*.md")):
        suite_id = path.name[:5]
        lines = path.read_text(encoding="utf-8").splitlines()
        current: dict | None = None

        for line in lines:
            match = CASE_RE.match(line)
            if match:
                if current:
                    cases.append(current)
                tc_id = match.group(1)
                if tc_id in seen:
                    print(f"  ! duplicate {tc_id} in {path.name} (also in {seen[tc_id]})")
                seen[tc_id] = path.name
                current = {
                    "id": tc_id, "title": clean(match.group(2)), "suite": suite_id,
                    "type": "", "priority": "", "screen": "", "persona": "", "automation": "",
                    "file": path.name,
                }
                continue

            if current is None:
                continue

            meta = META_RE.match(line)
            if meta:
                key = clean(meta.group(1)).lower()
                value = clean(meta.group(2))
                if key in ("field", "---") or set(value) <= {"-", ":"}:
                    continue
                if key == "suite":
                    current["suite_label"] = value
                elif key in ("type", "priority", "screen", "persona"):
                    current[key] = value
                elif key.startswith("automation"):
                    current["automation"] = value

        if current:
            cases.append(current)

    cases.sort(key=lambda c: (c["suite"], c["id"]))
    return cases


def read_existing(source: Path = OUT) -> tuple[dict[str, list], list[list], list[list]]:
    """Pull forward logged results, the Defects sheet, and the Daily Log."""
    if not source.exists():
        return {}, [], []

    book = load_workbook(source)
    results: dict[str, list] = {}
    if "Test Cases" in book.sheetnames:
        sheet = book["Test Cases"]
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if row and row[0]:
                # columns J..P hold everything a tester types
                results[str(row[0]).strip()] = list(row[9:16])

    def rows_of(name: str) -> list[list]:
        if name not in book.sheetnames:
            return []
        out = []
        for row in book[name].iter_rows(min_row=2, values_only=True):
            if any(cell not in (None, "") for cell in row):
                out.append(list(row))
        return out

    return results, rows_of("Defects"), rows_of("Daily Log")


def style_header(sheet, widths: list[int]) -> None:
    for idx, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(idx)].width = width
    for cell in sheet[1]:
        cell.font = HEAD
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[1].height = 28
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def build_instructions(book: Workbook, total: int, per_owner: dict[str, int]) -> None:
    sheet = book.create_sheet("How to use")
    sheet.column_dimensions["A"].width = 22
    sheet.column_dimensions["B"].width = 110
    sheet.sheet_view.showGridLines = False

    blocks = [
        ("Greenfield Test Execution Tracker", ""),
        ("", ""),
        ("What this is",
         "The execution log for the greenfield test cycle. The test cases themselves live in "
         "docs/qa/suites/ as Markdown — this workbook records what happened when you ran them."),
        ("Sheets",
         "Test Cases — one row per case, filter by Owner to find yours.  |  "
         "Defects — every failure you raise.  |  "
         "Daily Log — one row per tester per day.  |  "
         "Summary — live counts, no typing needed."),
        ("", ""),
        ("Every session", ""),
        ("1.",
         "Run the environment health check from GREENFIELD_TEST_PLAN.md §2 and record it on the Daily Log."),
        ("2.",
         "Filter Test Cases by your initial in the Owner column and work in TC-ID order within a suite."),
        ("3.",
         "Open the matching suite file in docs/qa/suites/ for the full steps — this sheet holds titles only."),
        ("4.",
         "Set Status for EVERY case you touch. Leaving it 'Not Run' means nobody looked, which is the one "
         "thing the tracker must never be ambiguous about."),
        ("5.",
         "On a failure: screenshot first, then add a row on Defects, then put the DEFECT-ID in the case row."),
        ("6.",
         "Blocked is not Fail. If a case cannot run because an earlier one failed, mark it Blocked and name "
         "the blocking DEFECT-ID — otherwise one root cause shows up as ten failures."),
        ("7.",
         "End of session: fill in the Daily Log and post the same one-liner in the team channel."),
        ("", ""),
        ("Re-testing",
         "When a defect is fixed, re-run the original case and fill the Re-test columns. A defect is closed "
         "by a passing test case, not by a developer's say-so."),
        ("Editing cases",
         "Never change a case's steps here. Edit the Markdown in docs/qa/suites/, then re-run "
         "'python docs/qa/build_tracker.py' — your logged results are preserved by TC-ID."),
        ("", ""),
        ("Cycle size", f"{total} test cases   ·   Tester A: {per_owner.get('A', 0)}   ·   "
                       f"Tester B: {per_owner.get('B', 0)}   ·   Tester C: {per_owner.get('C', 0)}"),
        ("Out of scope", "The brownfield track. Do not create a Brownfield project."),
    ]

    for row_idx, (label, body) in enumerate(blocks, start=1):
        sheet.cell(row=row_idx, column=1, value=label).alignment = TOP
        sheet.cell(row=row_idx, column=2, value=body).alignment = WRAP
        if row_idx == 1:
            sheet.cell(row=row_idx, column=1).font = TITLE
        elif label and not label[0].isdigit():
            sheet.cell(row=row_idx, column=1).font = Font(bold=True, color=NAVY)
        if body and len(body) > 95:
            sheet.row_dimensions[row_idx].height = 30


def build_cases(book: Workbook, cases: list[dict], carried: dict[str, list]) -> None:
    sheet = book.create_sheet("Test Cases")
    headers = [
        "TC-ID", "Suite", "Title", "Type", "Priority", "Screen", "Persona", "Automation",
        "Owner", "Status", "Run date", "Tester", "Actual result / notes", "Defect ID",
        "Re-test", "Re-test date",
    ]
    sheet.append(headers)

    for case in cases:
        owner = SUITE_OWNER.get(case["suite"], "")
        logged = carried.get(case["id"], [None] * 7)
        sheet.append([
            case["id"], case["suite"], case["title"], case["type"], case["priority"],
            case["screen"], case["persona"], case["automation"], owner,
            logged[0] or "Not Run", logged[1], logged[2], logged[3], logged[4],
            logged[5], logged[6],
        ])

    style_header(sheet, [10, 8, 62, 12, 10, 24, 14, 12, 8, 13, 12, 10, 46, 12, 12, 12])
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP
            cell.border = BOX

    last = sheet.max_row
    status_dv = DataValidation(type="list", formula1=f'"{",".join(STATUSES)}"', allow_blank=True)
    status_dv.error = "Pick one of: " + ", ".join(STATUSES)
    sheet.add_data_validation(status_dv)
    status_dv.add(f"J2:J{last}")

    retest_dv = DataValidation(type="list", formula1='"Not needed,Pending,Pass,Fail"', allow_blank=True)
    sheet.add_data_validation(retest_dv)
    retest_dv.add(f"O2:O{last}")

    for colour, formula in (
        ("FDECEA", 'INDIRECT("J"&ROW())="Fail"'),
        ("E9F7EF", 'INDIRECT("J"&ROW())="Pass"'),
        ("FFF6E5", 'INDIRECT("J"&ROW())="Blocked"'),
    ):
        from openpyxl.formatting.rule import FormulaRule
        sheet.conditional_formatting.add(
            f"A2:P{last}",
            FormulaRule(formula=[formula], fill=PatternFill("solid", fgColor=colour)),
        )


def build_defects(book: Workbook, rows: list[list]) -> None:
    sheet = book.create_sheet("Defects")
    headers = [
        "Defect ID", "Title", "Severity", "Status", "Found in TC", "Screen / route",
        "Steps to reproduce", "Expected", "Actual (verbatim)", "Evidence (screenshot path)",
        "Raised by", "Raised on", "Assigned to", "Fix ref / commit", "Re-tested by", "Closed on",
    ]
    sheet.append(headers)
    for row in rows:
        sheet.append(row[: len(headers)])

    style_header(sheet, [11, 52, 10, 22, 12, 22, 56, 40, 40, 26, 11, 12, 13, 18, 13, 12])
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP
            cell.border = BOX

    last = max(sheet.max_row, 200)
    sev_dv = DataValidation(type="list", formula1=f'"{",".join(SEVERITIES)}"', allow_blank=True)
    sheet.add_data_validation(sev_dv)
    sev_dv.add(f"C2:C{last}")
    state_dv = DataValidation(type="list", formula1=f'"{",".join(DEFECT_STATES)}"', allow_blank=True)
    sheet.add_data_validation(state_dv)
    state_dv.add(f"D2:D{last}")


def build_daily_log(book: Workbook, rows: list[list]) -> None:
    sheet = book.create_sheet("Daily Log")
    headers = [
        "Date", "Tester", "Hours", "Env health OK?", "Suites worked",
        "Cases run", "Pass", "Fail", "Blocked", "What got in the way",
    ]
    sheet.append(headers)
    for row in rows:
        sheet.append(row[: len(headers)])
    style_header(sheet, [12, 9, 8, 14, 26, 11, 8, 8, 10, 68])
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP
            cell.border = BOX

    last = max(sheet.max_row, 120)
    dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    sheet.add_data_validation(dv)
    dv.add(f"D2:D{last}")


def build_summary(book: Workbook, cases: list[dict]) -> None:
    sheet = book.create_sheet("Summary")
    sheet.sheet_view.showGridLines = False
    for col, width in zip("ABCDEFGH", (26, 12, 12, 12, 12, 12, 12, 14)):
        sheet.column_dimensions[col].width = width

    last = len(cases) + 1
    tc = f"'Test Cases'!"

    def header(row: int, text: str) -> None:
        cell = sheet.cell(row=row, column=1, value=text)
        cell.font = Font(bold=True, size=12, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        for col in range(2, 9):
            sheet.cell(row=row, column=col).fill = PatternFill("solid", fgColor=NAVY)

    sheet["A1"] = "Cycle Summary"
    sheet["A1"].font = TITLE
    sheet["A2"] = "Everything below is calculated. Do not type in this sheet."
    sheet["A2"].font = Font(italic=True, color="667085")

    header(4, "Overall")
    sheet.append([])
    row = 5
    sheet.cell(row=row, column=1, value="Total cases")
    sheet.cell(row=row, column=2, value=f"=COUNTA({tc}A2:A{last})")
    for offset, status in enumerate(STATUSES, start=1):
        sheet.cell(row=row + offset, column=1, value=status)
        sheet.cell(row=row + offset, column=2,
                   value=f'=COUNTIF({tc}J2:J{last},"{status}")')
        sheet.cell(row=row + offset, column=3,
                   value=f"=IFERROR(B{row + offset}/$B${row},0)").number_format = "0%"
    sheet.cell(row=row + 6, column=1, value="Executed (not 'Not Run')").font = Font(bold=True)
    sheet.cell(row=row + 6, column=2, value=f"=$B${row}-B{row + 1}").font = Font(bold=True)
    sheet.cell(row=row + 6, column=3, value=f"=IFERROR(B{row + 6}/$B${row},0)").number_format = "0%"

    base = row + 8
    header(base, "By owner")
    sheet.cell(row=base + 1, column=1, value="Owner").font = Font(bold=True)
    for idx, label in enumerate(["Total", "Pass", "Fail", "Blocked", "Not Run"], start=2):
        sheet.cell(row=base + 1, column=idx, value=label).font = Font(bold=True)
    for offset, owner in enumerate(["A", "B", "C"], start=2):
        r = base + offset
        sheet.cell(row=r, column=1, value=f"Tester {owner}")
        sheet.cell(row=r, column=2, value=f'=COUNTIF({tc}I2:I{last},"{owner}")')
        for col, status in zip(range(3, 7), ["Pass", "Fail", "Blocked", "Not Run"]):
            sheet.cell(row=r, column=col,
                       value=f'=COUNTIFS({tc}I2:I{last},"{owner}",{tc}J2:J{last},"{status}")')

    base += 6
    header(base, "By suite")
    sheet.cell(row=base + 1, column=1, value="Suite").font = Font(bold=True)
    for idx, label in enumerate(["Total", "Pass", "Fail", "Blocked", "Not Run", "% done"], start=2):
        sheet.cell(row=base + 1, column=idx, value=label).font = Font(bold=True)
    suites = sorted({c["suite"] for c in cases})
    for offset, suite in enumerate(suites, start=2):
        r = base + offset
        sheet.cell(row=r, column=1, value=suite)
        sheet.cell(row=r, column=2, value=f'=COUNTIF({tc}B2:B{last},"{suite}")')
        for col, status in zip(range(3, 7), ["Pass", "Fail", "Blocked", "Not Run"]):
            sheet.cell(row=r, column=col,
                       value=f'=COUNTIFS({tc}B2:B{last},"{suite}",{tc}J2:J{last},"{status}")')
        sheet.cell(row=r, column=7,
                   value=f"=IFERROR((B{r}-F{r})/B{r},0)").number_format = "0%"

    base += len(suites) + 3
    header(base, "By priority")
    sheet.cell(row=base + 1, column=1, value="Priority").font = Font(bold=True)
    for idx, label in enumerate(["Total", "Executed", "Fail"], start=2):
        sheet.cell(row=base + 1, column=idx, value=label).font = Font(bold=True)
    for offset, priority in enumerate(["Critical", "High", "Medium", "Low"], start=2):
        r = base + offset
        sheet.cell(row=r, column=1, value=priority)
        sheet.cell(row=r, column=2, value=f'=COUNTIF({tc}E2:E{last},"{priority}")')
        sheet.cell(row=r, column=3,
                   value=f'=B{r}-COUNTIFS({tc}E2:E{last},"{priority}",{tc}J2:J{last},"Not Run")')
        sheet.cell(row=r, column=4,
                   value=f'=COUNTIFS({tc}E2:E{last},"{priority}",{tc}J2:J{last},"Fail")')

    base += 7
    header(base, "Defects")
    for offset, severity in enumerate(SEVERITIES, start=1):
        r = base + offset
        sheet.cell(row=r, column=1, value=f"{severity} — open")
        sheet.cell(
            row=r, column=2,
            value=('=COUNTIFS(Defects!C2:C500,"%s",Defects!D2:D500,"<>Closed",'
                   'Defects!D2:D500,"<>Duplicate",Defects!D2:D500,"<>Won\'t Fix")' % severity),
        )
    sheet.cell(row=base + 5, column=1, value="Total raised")
    sheet.cell(row=base + 5, column=2, value="=COUNTA(Defects!A2:A500)")

    base += 7
    header(base, "Exit criteria")
    checks = [
        ("All Critical priority cases executed",
         f'=IF(COUNTIFS({tc}E2:E{last},"Critical",{tc}J2:J{last},"Not Run")=0,"MET","NOT MET")'),
        ("All High priority cases executed",
         f'=IF(COUNTIFS({tc}E2:E{last},"High",{tc}J2:J{last},"Not Run")=0,"MET","NOT MET")'),
        ("90% of all cases executed",
         f'=IF(($B${row}-B{row + 1})/$B${row}>=0.9,"MET","NOT MET")'),
        ("Zero open Critical defects", f'=IF(B{base - 6}=0,"MET","NOT MET")'),
        ("Zero open High defects", f'=IF(B{base - 5}=0,"MET","NOT MET")'),
    ]
    for offset, (label, formula) in enumerate(checks, start=1):
        sheet.cell(row=base + offset, column=1, value=label)
        cell = sheet.cell(row=base + offset, column=2, value=formula)
        cell.font = Font(bold=True)

    from openpyxl.formatting.rule import CellIsRule
    sheet.conditional_formatting.add(
        f"B{base + 1}:B{base + len(checks)}",
        CellIsRule(operator="equal", formula=['"MET"'],
                   fill=PatternFill("solid", fgColor="E9F7EF"), font=Font(color="1B7A47", bold=True)),
    )
    sheet.conditional_formatting.add(
        f"B{base + 1}:B{base + len(checks)}",
        CellIsRule(operator="equal", formula=['"NOT MET"'],
                   fill=PatternFill("solid", fgColor="FDECEA"), font=Font(color="B42318", bold=True)),
    )


def main() -> None:
    parser = argparse.ArgumentParser(description="Build the QA execution tracker.")
    parser.add_argument("--out", help="write a new workbook with this file name in docs/qa "
                                      "instead of overwriting the original")
    args = parser.parse_args()
    target = QA_DIR / args.out if args.out else OUT
    if target.parent != QA_DIR or target.suffix != ".xlsx":
        sys.exit("--out takes a plain .xlsx file name, written beside the original.")

    cases = parse_suites()
    if not cases:
        sys.exit("No test cases found — write the suite files first.")

    carried, defects, daily = read_existing(target if target.exists() else OUT)
    per_owner: dict[str, int] = {}
    for case in cases:
        owner = SUITE_OWNER.get(case["suite"], "?")
        per_owner[owner] = per_owner.get(owner, 0) + 1

    book = Workbook()
    book.remove(book.active)
    build_instructions(book, len(cases), per_owner)
    build_cases(book, cases, carried)
    build_defects(book, defects)
    build_daily_log(book, daily)
    build_summary(book, cases)
    book.save(target)

    print(f"Wrote {target.relative_to(QA_DIR.parent.parent)}")
    print(f"  {len(cases)} test cases across {len({c['suite'] for c in cases})} suites")
    for owner in sorted(per_owner):
        print(f"  Tester {owner}: {per_owner[owner]}")
    if carried:
        print(f"  carried forward results for {len(carried)} case(s)")
    if defects or daily:
        print(f"  preserved {len(defects)} defect row(s), {len(daily)} daily-log row(s)")


if __name__ == "__main__":
    main()
